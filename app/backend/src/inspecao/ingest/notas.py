"""Leitura da planilha de notas de ocorrencia e upsert em notas_ocorrencia.

Planilha: uma aba, 62 linhas (1 cabecalho + 61 codigos), colunas
NOTA | DESCRICAO | NOTA EXIGE FOTO?  com valores SIM ou NAO.
"""

import logging
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session

from inspecao.models import NotaOcorrencia

log = logging.getLogger(__name__)

# Cena esperada por codigo. So os codigos com volume real nos dados receberam
# mapeamento. Ver doc 02 do plano, secao 6.1: todos esperam o display, e e por
# isso que o classificador de cena nao e treinavel com a base atual.
CENA_ESPERADA = {
    "T181": "display_medidor",
    "P111": "display_medidor",
    "L101": "display_medidor",
    "T111": "display_medidor",
    "B111": "display_medidor",
    "L131": "display_medidor",
    "T161": "display_medidor",
    "M141": "display_medidor",
}


@dataclass(slots=True)
class NotaPlanilha:
    codigo: str
    descricao: str
    exige_foto: bool


def ler_planilha(caminho: Path) -> list[NotaPlanilha]:
    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        aba = livro[livro.sheetnames[0]]
        notas: list[NotaPlanilha] = []
        for codigo, descricao, exige in aba.iter_rows(min_row=2, max_col=3, values_only=True):
            if not codigo:
                continue
            notas.append(
                NotaPlanilha(
                    codigo=str(codigo).strip(),
                    descricao=str(descricao or "").strip(),
                    exige_foto=str(exige or "").strip().upper() == "SIM",
                )
            )
        return notas
    finally:
        livro.close()


def sincronizar(sessao: Session, caminho: Path) -> int:
    """Faz upsert das notas. Uma planilha atualizada entra sem migracao."""
    notas = ler_planilha(caminho)
    if not notas:
        log.warning("planilha de notas vazia: %s", caminho)
        return 0

    linhas = [
        {
            "codigo": n.codigo,
            "descricao": n.descricao,
            "exige_foto": n.exige_foto,
            "cena_esperada": CENA_ESPERADA.get(n.codigo),
        }
        for n in notas
    ]

    stmt = insert(NotaOcorrencia).values(linhas)
    stmt = stmt.on_conflict_do_update(
        index_elements=["codigo"],
        set_={
            "descricao": stmt.excluded.descricao,
            "exige_foto": stmt.excluded.exige_foto,
            "cena_esperada": stmt.excluded.cena_esperada,
        },
    )
    sessao.execute(stmt)
    log.info("notas de ocorrencia sincronizadas: %d", len(linhas))
    return len(linhas)
